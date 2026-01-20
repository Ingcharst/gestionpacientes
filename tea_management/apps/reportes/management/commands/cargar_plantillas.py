"""
Comando de Django para cargar plantillas de informes desde archivos Word.

Uso:
    python manage.py cargar_plantillas
"""

from django.core.management.base import BaseCommand
from django.conf import settings
from pathlib import Path
import sys


class Command(BaseCommand):
    help = 'Carga plantillas de informes desde archivos Word (.docx)'

    def add_arguments(self, parser):
        parser.add_argument(
            '--dir',
            type=str,
            default='plantillas_informes',
            help='Directorio donde están los archivos Word (default: plantillas_informes/)'
        )

    def handle(self, *args, **options):
        self.stdout.write("="*70)
        self.stdout.write(self.style.SUCCESS("CARGA DE PLANTILLAS DE INFORMES"))
        self.stdout.write("="*70)

        # ================================================================
        # 1. VERIFICAR LIBRERÍA PYTHON-DOCX
        # ================================================================
        self.stdout.write("\n[1/6] Verificando python-docx...")
        try:
            from docx import Document
            self.stdout.write(self.style.SUCCESS("   ✅ python-docx instalado"))
        except ImportError:
            self.stdout.write(self.style.ERROR("   ❌ ERROR: Falta python-docx"))
            self.stdout.write("   Ejecuta: pip install python-docx")
            return

        # ================================================================
        # 2. VERIFICAR MODELO PlantillaInforme
        # ================================================================
        self.stdout.write("\n[2/6] Verificando modelo PlantillaInforme...")
        try:
            from apps.reportes.models import PlantillaInforme
            self.stdout.write(self.style.SUCCESS("   ✅ Modelo PlantillaInforme encontrado"))
        except ImportError as e:
            self.stdout.write(self.style.ERROR(f"   ❌ ERROR: {e}"))
            self.stdout.write("   Verifica que existe apps/reportes/models.py")
            return

        # ================================================================
        # 3. VERIFICAR MODELO Terapia
        # ================================================================
        self.stdout.write("\n[3/6] Verificando modelo Terapia...")
        try:
            from apps.terapias.models import Terapia
            self.stdout.write(self.style.SUCCESS("   ✅ Modelo Terapia encontrado"))
        except ImportError as e:
            self.stdout.write(self.style.ERROR(f"   ❌ ERROR: {e}"))
            return

        # ================================================================
        # 4. BUSCAR TERAPIA COGNITIVA
        # ================================================================
        self.stdout.write("\n[4/6] Buscando terapia 'Cognitiva'...")
        terapia_cognitiva = Terapia.objects.filter(nombre__icontains='cognitiva').first()
        
        if not terapia_cognitiva:
            self.stdout.write(self.style.ERROR("   ❌ No existe terapia con 'cognitiva' en el nombre"))
            self.stdout.write("\n   Solución: Crear la terapia primero:")
            self.stdout.write("   python manage.py shell")
            self.stdout.write("   >>> from apps.terapias.models import Terapia")
            self.stdout.write("   >>> Terapia.objects.create(")
            self.stdout.write("   ...     codigo='TC-001',")
            self.stdout.write("   ...     nombre='Terapia Cognitiva',")
            self.stdout.write("   ...     especialidad='COGNITIVO',")
            self.stdout.write("   ...     modalidad='INDIVIDUAL',")
            self.stdout.write("   ...     activo=True")
            self.stdout.write("   ... )")
            return
        
        self.stdout.write(self.style.SUCCESS(f"   ✅ Terapia: {terapia_cognitiva.nombre} ({terapia_cognitiva.codigo})"))

        # ================================================================
        # 5. CONFIGURAR DIRECTORIOS
        # ================================================================
        self.stdout.write("\n[5/6] Configurando directorios...")
        
        # Directorio base del proyecto (donde está manage.py)
        BASE_DIR = Path(settings.BASE_DIR)
        self.stdout.write(f"   Directorio base: {BASE_DIR}")
        
        # Directorio de plantillas
        plantillas_dir_name = options['dir']
        PLANTILLAS_DIR = BASE_DIR / plantillas_dir_name
        self.stdout.write(f"   Directorio plantillas: {PLANTILLAS_DIR}")
        
        # Verificar que existe
        if not PLANTILLAS_DIR.exists():
            self.stdout.write(self.style.WARNING(f"   ⚠️  No existe: {PLANTILLAS_DIR}"))
            self.stdout.write("   Creando directorio...")
            PLANTILLAS_DIR.mkdir(parents=True, exist_ok=True)
            self.stdout.write(self.style.SUCCESS(f"   ✅ Directorio creado"))
            self.stdout.write("\n" + "="*70)
            self.stdout.write("📋 ACCIÓN REQUERIDA:")
            self.stdout.write(f"   Coloca los archivos Word en: {PLANTILLAS_DIR}")
            self.stdout.write("   Archivos esperados:")
            self.stdout.write("   - 3-6_AÑOS.docx")
            self.stdout.write("   - 7-11_AÑOS.docx")
            self.stdout.write("   - 12-16_AÑOS.docx")
            self.stdout.write("="*70)
            return
        
        self.stdout.write(self.style.SUCCESS(f"   ✅ Directorio existe"))

        # ================================================================
        # 6. PROCESAR PLANTILLAS
        # ================================================================
        self.stdout.write("\n[6/6] Procesando plantillas...")
        
        # Función para extraer texto
        def extraer_texto_completo(doc_path):
            """Extrae todo el texto del documento Word"""
            try:
                doc = Document(doc_path)
                parrafos = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        parrafos.append(para.text.strip())
                return '\n\n'.join(parrafos)
            except Exception as e:
                self.stdout.write(self.style.ERROR(f"      ❌ Error leyendo {doc_path.name}: {e}"))
                return None

        # Configuración de plantillas
        plantillas_config = [
            {
                'archivo': PLANTILLAS_DIR / '3-6_AÑOS.docx',
                'rango': '3-6',
                'nombre': '3-6 años',
                'terapia': terapia_cognitiva
            },
            {
                'archivo': PLANTILLAS_DIR / '7-11_AÑOS.docx',
                'rango': '7-11',
                'nombre': '7-11 años',
                'terapia': terapia_cognitiva
            },
            {
                'archivo': PLANTILLAS_DIR / '12-16_AÑOS.docx',
                'rango': '12-16',
                'nombre': '12-16 años',
                'terapia': terapia_cognitiva
            }
        ]

        procesadas = 0
        errores = 0

        for i, config in enumerate(plantillas_config, 1):
            archivo = config['archivo']
            self.stdout.write(f"\n   [{i}/{len(plantillas_config)}] {archivo.name}")
            
            # Verificar archivo existe
            if not archivo.exists():
                self.stdout.write(self.style.WARNING(f"      ⚠️  No existe, saltando..."))
                errores += 1
                continue
            
            # Extraer contenido
            contenido = extraer_texto_completo(archivo)
            if contenido is None:
                errores += 1
                continue
            
            self.stdout.write(f"      📄 Extraído: {len(contenido)} caracteres")
            
            # Reemplazar placeholders
            contenido = contenido.replace('XXXXXXXX', '{nombre_paciente}')
            contenido = contenido.replace('xx edad', '{edad}')
            contenido = contenido.replace('XX EDAD', '{edad}')
            contenido = contenido.replace('xx años', '{edad} años')
            
            # Separar secciones
            partes = contenido.split('Familia:')
            if len(partes) > 1:
                parte_familia = partes[1].split('Escolaridad')
                recomendaciones_familia = parte_familia[0].strip() if parte_familia else ''
                recomendaciones_escuela = parte_familia[1].strip() if len(parte_familia) > 1 else ''
                contenido_principal = partes[0].strip()
            else:
                contenido_principal = contenido
                recomendaciones_familia = ''
                recomendaciones_escuela = ''
            
            # Guardar en BD
            try:
                plantilla, created = PlantillaInforme.objects.update_or_create(
                    terapia=config['terapia'],
                    rango_edad=config['rango'],
                    defaults={
                        'titulo': 'INFORME DE TERAPIAS COGNITIVAS',
                        'contenido_principal': contenido_principal,
                        'recomendaciones_familia': recomendaciones_familia,
                        'recomendaciones_escuela': recomendaciones_escuela,
                        'activo': True
                    }
                )
                
                accion = "Creada" if created else "Actualizada"
                self.stdout.write(self.style.SUCCESS(f"      ✅ {accion}: {config['nombre']}"))
                procesadas += 1
                
            except Exception as e:
                self.stdout.write(self.style.ERROR(f"      ❌ Error guardando: {e}"))
                errores += 1

        # ================================================================
        # RESUMEN
        # ================================================================
        self.stdout.write("\n" + "="*70)
        self.stdout.write("RESUMEN")
        self.stdout.write("="*70)
        self.stdout.write(f"✅ Procesadas exitosamente: {procesadas}")
        if errores > 0:
            self.stdout.write(f"⚠️  Con errores: {errores}")
        
        if procesadas > 0:
            self.stdout.write("\n" + self.style.SUCCESS("🎉 ¡Plantillas cargadas exitosamente!"))
            self.stdout.write("\nVerificar:")
            self.stdout.write("python manage.py shell")
            self.stdout.write(">>> from apps.reportes.models import PlantillaInforme")
            self.stdout.write(">>> PlantillaInforme.objects.all()")
        else:
            self.stdout.write(self.style.WARNING("\n⚠️  No se procesó ninguna plantilla"))
            self.stdout.write(f"Coloca los archivos Word en: {PLANTILLAS_DIR}")
        
        self.stdout.write("="*70)