"""
Utilidades para generación de reportes en PDF y Excel.
"""
from io import BytesIO
from datetime import datetime, date
from decimal import Decimal
from tkinter import Image
from django.core.files.base import ContentFile
from django.http import HttpResponse
from django.db.models import Sum, Count, Q
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT, TA_JUSTIFY
import os

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    Workbook = None


class ReportePDFGenerator:
    """Generador base para reportes en PDF."""
    
    def __init__(self, titulo, subtitulo=None):
        self.titulo = titulo
        self.subtitulo = subtitulo
        self.buffer = BytesIO()
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()
    
    def _setup_custom_styles(self):
        """Configurar estilos personalizados."""
        self.styles.add(ParagraphStyle(
            name='TituloReporte',
            parent=self.styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1a5490'),
            spaceAfter=12,
            alignment=TA_CENTER
        ))
        self.styles.add(ParagraphStyle(
            name='SubtituloReporte',
            parent=self.styles['Normal'],
            fontSize=12,
            textColor=colors.grey,
            spaceAfter=20,
            alignment=TA_CENTER
        ))
    
    def crear_documento(self):
        """Crear documento PDF."""
        doc = SimpleDocTemplate(
            self.buffer,
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=18,
        )
        return doc
    
    def crear_encabezado(self):
        """Crear encabezado del reporte."""
        elementos = []
        elementos.append(Paragraph(self.titulo, self.styles['TituloReporte']))
        if self.subtitulo:
            elementos.append(Paragraph(self.subtitulo, self.styles['SubtituloReporte']))
        elementos.append(Spacer(1, 0.2*inch))
        return elementos
    
    def crear_tabla(self, datos, col_widths=None, header_color=None):
        """Crear tabla formateada."""
        if not header_color:
            header_color = colors.HexColor('#4e73df')
        
        tabla = Table(datos, colWidths=col_widths)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), header_color),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        return tabla
    
    def get_response(self, filename):
        """Obtener HttpResponse con el PDF."""
        self.buffer.seek(0)
        response = HttpResponse(self.buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class ReporteExcelGenerator:
    """Generador base para reportes en Excel."""
    
    def __init__(self, titulo):
        if not Workbook:
            raise ImportError("openpyxl no está instalado")
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = titulo[:31]  # Excel limita a 31 caracteres
        self.current_row = 1
    
    def agregar_titulo(self, titulo, columnas=5):
        """Agregar título principal."""
        self.ws.merge_cells(start_row=self.current_row, start_column=1,
                            end_row=self.current_row, end_column=columnas)
        cell = self.ws.cell(row=self.current_row, column=1, value=titulo)
        cell.font = Font(size=16, bold=True, color="1F4E78")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        self.current_row += 2
    
    def agregar_encabezados(self, headers):
        """Agregar fila de encabezados."""
        for col, header in enumerate(headers, start=1):
            cell = self.ws.cell(row=self.current_row, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        self.current_row += 1
    
    def agregar_fila(self, datos):
        """Agregar fila de datos."""
        for col, dato in enumerate(datos, start=1):
            cell = self.ws.cell(row=self.current_row, column=col, value=dato)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        self.current_row += 1
    
    def ajustar_columnas(self):
        """Ajustar ancho de columnas automáticamente."""
        for column in self.ws.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            self.ws.column_dimensions[column[0].column_letter].width = adjusted_width
    
    def get_response(self, filename):
        """Obtener HttpResponse con el Excel."""
        buffer = BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


def calcular_totales_procedimientos(procedimientos):
    """Calcular totales de procedimientos."""
    total = procedimientos.aggregate(
        total_monto=Sum('costo')
    )['total_monto'] or Decimal('0.00')
    
    cantidad = procedimientos.count()
    
    return {
        'cantidad': cantidad,
        'total': total
    }


def calcular_totales_sesiones(sesiones):
    """Calcular totales de sesiones terapéuticas."""
    total = sesiones.aggregate(
        total_monto=Sum('costo')
    )['total_monto'] or Decimal('0.00')
    
    cantidad = sesiones.count()
    
    return {
        'cantidad': cantidad,
        'total': total
    }


def formato_moneda(valor):
    """Formatear valor como moneda colombiana."""
    if valor is None:
        return "$0"
    return f"${valor:,.0f}".replace(",", ".")


def obtener_nombre_mes(numero_mes):
    """Obtener nombre del mes en español."""
    meses = {
        1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
        5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
        9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
    }
    return meses.get(numero_mes, '')


def generar_pdf_informe(informe):
    """
    Genera un PDF del informe de evolución con tabla de identificación y firma.
    
    Args:
        informe: Instancia de InformeEvolucion
    
    Returns:
        ContentFile con el PDF generado
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter,
                            topMargin=0.5*inch, bottomMargin=0.5*inch,
                            leftMargin=0.75*inch, rightMargin=0.75*inch)
    
    # Estilos
    styles = getSampleStyleSheet()
    
    titulo_style = ParagraphStyle(
        'TituloInforme',
        parent=styles['Heading1'],
        fontSize=14,
        textColor=colors.HexColor('#1a5490'),
        spaceAfter=20,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )
    
    normal_style = ParagraphStyle(
        'Normal',
        parent=styles['Normal'],
        fontSize=10,
        alignment=TA_JUSTIFY,
        spaceAfter=10
    )
    
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Heading2'],
        fontSize=11,
        textColor=colors.HexColor('#2c5f8d'),
        spaceAfter=10,
        fontName='Helvetica-Bold'
    )
    
    # Elementos del documento
    elementos = []
    
    # 1. TÍTULO
    titulo = Paragraph(
        informe.plantilla_usada.titulo if informe.plantilla_usada else 'INFORME DE EVOLUCIÓN',
        titulo_style
    )
    elementos.append(titulo)
    elementos.append(Spacer(1, 0.2*inch))
    
    # 2. TABLA DE IDENTIFICACIÓN
    paciente = informe.paciente
    admision = informe.admision
    
    datos_tabla = [
        ['Nombres y Apellidos:', paciente.nombre_completo, 
            'Documento de identidad:', paciente.numero_documento],
        ['Edad:', f"{int(paciente.edad_actual)} años" if paciente.edad_actual else 'N/A',
            'Fecha de Nacimiento:', paciente.fecha_nacimiento.strftime('%d/%m/%Y')],
        ['Admisión:', admision.numero_admision,
            'Número de Intervenciones:', str(admision.cantidad_realizada)],
        ['Escolaridad:', getattr(paciente, 'nivel_escolar', 'N/A'),
            'Acudiente:', paciente.nombre_responsable],
        ['Diagnóstico:', paciente.diagnostico_principal,
            'Profesional:', informe.profesional.get_full_name()],
        ['EPS:', paciente.eps or 'N/A',
            'Autorización:', admision.numero_admision]
    ]
    
    tabla_identificacion = Table(datos_tabla, colWidths=[1.5*inch, 2*inch, 1.5*inch, 2*inch])
    tabla_identificacion.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#e8f4f8')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#e8f4f8')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    
    elementos.append(tabla_identificacion)
    elementos.append(Spacer(1, 0.3*inch))
    
    # 3. CONTENIDO DEL INFORME
    # Dividir contenido en párrafos
    contenido_texto = informe.contenido_informe
    parrafos = contenido_texto.split('\n\n')
    
    for parrafo in parrafos:
        if parrafo.strip():
            # Detectar si es un título (todo mayúsculas o termina en :)
            if parrafo.isupper() or parrafo.strip().endswith(':'):
                p = Paragraph(parrafo, subtitle_style)
            else:
                p = Paragraph(parrafo, normal_style)
            elementos.append(p)
            elementos.append(Spacer(1, 0.1*inch))
    
    # 4. OBSERVACIONES ESPECÍFICAS
    if informe.observaciones_especificas:
        elementos.append(Spacer(1, 0.2*inch))
        elementos.append(Paragraph('OBSERVACIONES ADICIONALES DEL PROFESIONAL', subtitle_style))
        elementos.append(Paragraph(informe.observaciones_especificas, normal_style))
    
    # 5. RECOMENDACIONES (si están en la plantilla)
    if informe.plantilla_usada:
        if informe.plantilla_usada.recomendaciones_familia:
            elementos.append(Spacer(1, 0.2*inch))
            elementos.append(Paragraph('RECOMENDACIONES FAMILIA', subtitle_style))
            elementos.append(Paragraph(informe.plantilla_usada.recomendaciones_familia, normal_style))
        
        if informe.plantilla_usada.recomendaciones_escuela:
            elementos.append(Spacer(1, 0.2*inch))
            elementos.append(Paragraph('RECOMENDACIONES ESCOLARIDAD', subtitle_style))
            elementos.append(Paragraph(informe.plantilla_usada.recomendaciones_escuela, normal_style))
    
    # 6. FIRMA DEL PROFESIONAL
    elementos.append(Spacer(1, 0.5*inch))
    
    if informe.firma_profesional and os.path.exists(informe.firma_profesional.path):
        try:
            # Insertar imagen de firma
            firma_img = Image(informe.firma_profesional.path, width=2*inch, height=1*inch)
            firma_img.hAlign = 'CENTER'
            elementos.append(firma_img)
        except:
            # Si falla, mostrar texto
            elementos.append(Paragraph('_________________________', normal_style))
    else:
        elementos.append(Paragraph('_________________________', normal_style))
    
    # Nombre del profesional
    firma_texto = Paragraph(
        f'<b>{informe.profesional.get_full_name()}</b><br/>'
        f'{informe.profesional.get_rol_display()}<br/>'
        f'Fecha: {informe.fecha_generacion.strftime("%d/%m/%Y")}',
        ParagraphStyle('Firma', parent=styles['Normal'], alignment=TA_CENTER, fontSize=9)
    )
    elementos.append(firma_texto)
    
    # Construir PDF
    doc.build(elementos)
    
    # Guardar en ContentFile
    buffer.seek(0)
    filename = f'informe_{informe.paciente.numero_documento}_{informe.fecha_generacion.strftime("%Y%m%d")}.pdf'
    
    return ContentFile(buffer.read(), name=filename)




