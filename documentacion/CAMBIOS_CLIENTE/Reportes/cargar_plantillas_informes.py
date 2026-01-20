"""
Script para cargar plantillas de informes desde archivos Word.
Ejecutar: python manage.py shell < cargar_plantillas_informes.py
"""

from docx import Document
from apps.reportes.models import PlantillaInforme
from apps.terapias.models import Terapia

def extraer_texto_completo(doc_path):
    """Extrae todo el texto del documento Word"""
    doc = Document(doc_path)
    parrafos = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            parrafos.append(para.text.strip())
    
    return '\n\n'.join(parrafos)

# Obtener terapia Cognitiva (ajustar según tu BD)
try:
    terapia_cognitiva = Terapia.objects.get(nombre__icontains='cognitiva')
except Terapia.DoesNotExist:
    print("ERROR: No existe terapia 'Cognitiva'. Créala primero.")
    exit(1)

# Cargar plantillas
plantillas = [
    {
        'archivo': '/mnt/user-data/uploads/3-6_AÑOS.docx',
        'rango': '3-6',
        'terapia': terapia_cognitiva
    },
    {
        'archivo': '/mnt/user-data/uploads/7-11_AÑOS_.docx',
        'rango': '7-11',
        'terapia': terapia_cognitiva
    },
    {
        'archivo': '/mnt/user-data/uploads/12-16_AÑOS_.docx',
        'rango': '12-16',
        'terapia': terapia_cognitiva
    }
]

for config in plantillas:
    contenido = extraer_texto_completo(config['archivo'])
    
    # Reemplazar placeholders genéricos con variables
    contenido = contenido.replace('XXXXXXXX', '{nombre_paciente}')
    contenido = contenido.replace('xx edad', '{edad}')
    
    # Separar recomendaciones (texto después de "Familia:" y "Escolaridad")
    partes = contenido.split('Familia:')
    if len(partes) > 1:
        parte_familia = partes[1].split('Escolaridad')
        recomendaciones_familia = parte_familia[0].strip() if parte_familia else ''
        recomendaciones_escuela = parte_familia[1].strip() if len(parte_familia) > 1 else ''
        contenido_principal = partes[0].strip()
    else:
        contenido_principal = contenido
        recomendaciones_familia = ''
        recomendaciones_escuela = ''
    
    # Crear o actualizar plantilla
    plantilla, created = PlantillaInforme.objects.update_or_create(
        terapia=config['terapia'],
        rango_edad=config['rango'],
        defaults={
            'titulo': 'INFORME DE TERAPIAS COGNITIVAS',
            'contenido_principal': contenido_principal,
            'recomendaciones_familia': recomendaciones_familia,
            'recomendaciones_escuela': recomendaciones_escuela,
            'activo': True
        }
    )
    
    accion = "Creada" if created else "Actualizada"
    print(f"{accion} plantilla: {config['terapia'].nombre} - {config['rango']} años")

print("\n✅ Plantillas cargadas exitosamente")
